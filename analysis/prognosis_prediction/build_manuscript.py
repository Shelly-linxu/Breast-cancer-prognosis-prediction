#!/usr/bin/env python3

from pathlib import Path
import re
import pandas as pd
from PIL import Image, ImageDraw, ImageFont

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / "analysis" / "prognosis_prediction" / "results"
PAPER = RESULTS / "paper_outputs"
OUTPUT = PAPER / "GBCS_five_year_prognosis_model_manuscript_WITH_STUDY_FLOW.docx"
FLOW_FIGURE = PAPER / "figure1_study_flow.png"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_GRAY = "F4F6F9"
MID_GRAY = "D9E1E8"
TEXT_GRAY = "666666"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C2CC", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(9)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def font_run(run, size=None, bold=None, italic=None, color=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc, text, bold_prefix=None, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             before=0, after=8, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.333
    p.paragraph_format.keep_with_next = keep_with_next
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        font_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        font_run(r2, italic=italic)
    else:
        r = p.add_run(text)
        font_run(r, italic=italic)
    return p


def add_caption(doc, label, text, above=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8 if above else 4)
    p.paragraph_format.space_after = Pt(4 if above else 10)
    p.paragraph_format.keep_with_next = above
    r = p.add_run(f"{label}. ")
    font_run(r, size=10, bold=True, color=DARK_BLUE)
    r = p.add_run(text)
    font_run(r, size=10)
    return p


def add_table(doc, headers, rows, widths_dxa, numeric_cols=None, font_size=9):
    numeric_cols = set(numeric_cols or [])
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for j, value in enumerate(headers):
        cell = hdr.cells[j]
        set_cell_shading(cell, LIGHT_GRAY)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(str(value))
        font_run(r, size=font_size, bold=True, color=DARK_BLUE)
    for row_data in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row_data):
            cell = cells[j]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in numeric_cols else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(value))
            font_run(r, size=font_size)
    set_table_geometry(table, widths_dxa)
    return table


def add_figure(doc, path, width_inches, caption_label, caption_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))
    add_caption(doc, caption_label, caption_text, above=False)


def build_study_flow_figure(path):
    """Create a journal-ready study flow diagram from verified analysis counts."""
    width, height = 1584, 1837
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    blue = "#2E74B5"
    dark_blue = "#1F4D78"
    pale_blue = "#EAF2F8"
    pale_gray = "#F3F5F7"
    border_gray = "#AAB5BF"
    text = "#263238"

    font_candidates = [
        "/System/Library/Fonts/Supplemental/Arial.ttf",
        "/Library/Fonts/Arial.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    bold_candidates = [
        "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        "/Library/Fonts/Arial Bold.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
    ]
    font_path = next((p for p in font_candidates if Path(p).exists()), font_candidates[-1])
    bold_path = next((p for p in bold_candidates if Path(p).exists()), bold_candidates[-1])

    def px(x):
        return int(round(x * width))

    def py(y):
        return int(round((1 - y) * height))

    def font(size, weight="normal"):
        return ImageFont.truetype(bold_path if weight == "bold" else font_path, int(round(size * 3.05)))

    def box(x, y, w, h, label, fill=pale_blue, edge=blue, fontsize=8.7,
            weight="normal", linestyle="solid"):
        left, top, right, bottom = px(x), py(y + h), px(x + w), py(y)
        if linestyle == "dashed":
            draw.rounded_rectangle((left, top, right, bottom), radius=16, fill=fill)
            dash, gap = 18, 10
            for sx in range(left + 16, right - 16, dash + gap):
                draw.line((sx, top, min(sx + dash, right - 16), top), fill=edge, width=3)
                draw.line((sx, bottom, min(sx + dash, right - 16), bottom), fill=edge, width=3)
            for sy in range(top + 16, bottom - 16, dash + gap):
                draw.line((left, sy, left, min(sy + dash, bottom - 16)), fill=edge, width=3)
                draw.line((right, sy, right, min(sy + dash, bottom - 16)), fill=edge, width=3)
        else:
            draw.rounded_rectangle((left, top, right, bottom), radius=16, fill=fill, outline=edge, width=3)
        draw.multiline_text(
            ((left + right) / 2, (top + bottom) / 2), label,
            font=font(fontsize, weight), fill=text, anchor="mm", align="center", spacing=7,
        )

    def arrow(x1, y1, x2, y2, color=dark_blue, linestyle="solid"):
        start, end = (px(x1), py(y1)), (px(x2), py(y2))
        draw.line((start, end), fill=color, width=3)
        ex, ey = end
        if abs(x2 - x1) >= abs(y2 - y1):
            direction = 1 if x2 > x1 else -1
            points = [(ex, ey), (ex - direction * 13, ey - 8), (ex - direction * 13, ey + 8)]
        else:
            direction = 1 if py(y2) > py(y1) else -1
            points = [(ex, ey), (ex - 8, ey - direction * 13), (ex + 8, ey - direction * 13)]
        draw.polygon(points, fill=color)

    def branch(x, y_top, y_bottom, x_targets):
        draw.line((px(x), py(y_top), px(x), py(y_bottom)), fill=dark_blue, width=3)
        draw.line((px(min(x_targets)), py(y_bottom), px(max(x_targets)), py(y_bottom)), fill=dark_blue, width=3)
        for target in x_targets:
            arrow(target, y_bottom, target, y_bottom - 0.025)

    # Participant selection: all counts were reproduced directly from the source file.
    box(0.19, 0.905, 0.62, 0.065,
        "GBCS linked clinical and 2023 follow-up database\nParticipants with breast-cancer records, n = 5,412",
        weight="bold", fontsize=9.2)
    arrow(0.50, 0.905, 0.50, 0.855)

    box(0.08, 0.775, 0.54, 0.075,
        "Invasive breast cancer\nClassification code indicating invasive disease, n = 5,057")
    box(0.69, 0.775, 0.29, 0.075,
        "Excluded, n = 355\nNon-invasive or other classification",
        fill=pale_gray, edge=border_gray, fontsize=8.2)
    arrow(0.50, 0.855, 0.50, 0.850)
    arrow(0.62, 0.813, 0.69, 0.813, color=border_gray)
    arrow(0.35, 0.775, 0.35, 0.730)

    box(0.08, 0.650, 0.54, 0.075,
        "Eligible diagnosis period\n1 October 2008 to 31 January 2018, n = 4,356")
    box(0.69, 0.650, 0.29, 0.075,
        "Excluded, n = 701\nMissing or outside diagnosis window",
        fill=pale_gray, edge=border_gray, fontsize=8.2)
    arrow(0.62, 0.688, 0.69, 0.688, color=border_gray)
    arrow(0.35, 0.650, 0.35, 0.605)

    box(0.08, 0.515, 0.54, 0.085,
        "Final analysis cohort, n = 4,231\nValid nonnegative OS and PFS times and event indicators\n492 deaths during follow-up; 293 deaths within five years",
        fill="#DDEBF7", weight="bold", fontsize=8.8)
    box(0.69, 0.520, 0.29, 0.075,
        "Excluded, n = 125\nMissing or invalid survival follow-up",
        fill=pale_gray, edge=border_gray, fontsize=8.2)
    arrow(0.62, 0.558, 0.69, 0.558, color=border_gray)

    # Analysis sets and validation pathways.
    branch(0.35, 0.515, 0.475, [0.235, 0.765])
    box(0.02, 0.325, 0.43, 0.125,
        "Model development on the full cohort\nn = 4,231; 10 multiply imputed datasets\nPrimary Cox model: age, stage, ER, PR,\nHER2, and Ki-67",
        fontsize=8.4)
    box(0.55, 0.325, 0.43, 0.125,
        "Temporal validation\nDevelopment: diagnosis through 2014, n = 2,577\nValidation: diagnosis from 2015, n = 1,654",
        fontsize=8.4)

    branch(0.235, 0.325, 0.285, [0.165, 0.50])
    arrow(0.765, 0.325, 0.765, 0.285)
    box(0.01, 0.135, 0.31, 0.125,
        "Bootstrap internal validation\n50 samples in each imputation\n500 replicates per model and endpoint\nOptimism-corrected performance",
        fill="#EEF5FA", fontsize=8.1)
    box(0.345, 0.135, 0.31, 0.125,
        "Internal-external validation\nLeave one GBCS hospital out in turn\nThree hospital hold-outs\nSupporting analysis",
        fill="#EEF5FA", fontsize=8.1)
    box(0.68, 0.135, 0.31, 0.125,
        "PREDICT v2.2 benchmark\nEligible temporal-validation subset, n = 1,371\nNot eligible for benchmark, n = 283\nPaired discrimination and calibration",
        fill="#EEF5FA", fontsize=7.8)

    box(0.10, 0.025, 0.80, 0.065,
        "Independent external validation in a non-GBCS healthcare system: not yet performed",
        fill="white", edge=border_gray, fontsize=8.5, weight="bold", linestyle="dashed")

    image.save(path, dpi=(220, 220), optimize=True)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(section, first=False):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = first

    if not first:
        hp = section.header.paragraphs[0]
        hp.text = "GBCS five-year prognosis model - manuscript draft"
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after = Pt(0)
        for run in hp.runs:
            font_run(run, size=9, color=TEXT_GRAY)
    fp = section.footer.paragraphs[0]
    add_page_number(fp)


def h(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def fmt(value, digits=3):
    return f"{float(value):.{digits}f}"


def build_document():
    build_study_flow_figure(FLOW_FIGURE)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], first=True)

    # First-page editorial title block, simplified for an academic manuscript.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("ORIGINAL RESEARCH | PROGNOSTIC MODEL DEVELOPMENT")
    font_run(r, size=10, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Development and internal validation of a diagnosis-time model for five-year mortality after invasive breast cancer in South China: the Guangzhou Breast Cancer Study")
    font_run(r, size=18, bold=True, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("Author names, order, degrees, and affiliations to be confirmed")
    font_run(r, size=11, italic=True, color=TEXT_GRAY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Corresponding author: to be confirmed | Draft updated 19 July 2026")
    font_run(r, size=9.5, color=TEXT_GRAY)

    h(doc, "Abstract", 1)
    abstract_parts = [
        ("Background: ", "Breast cancer prognosis models developed in Western populations may not provide well-calibrated absolute risks for patients in China. We developed and internally validated a parsimonious diagnosis-time model for five-year mortality in the Guangzhou Breast Cancer Study (GBCS)."),
        ("Methods: ", "Women with invasive breast cancer diagnosed from October 2008 to January 2018 and valid survival follow-up were included. The primary Cox model used age, stage, estrogen receptor, progesterone receptor, HER2, and Ki-67. Missing predictors were handled using multiple imputation. Performance was evaluated using the C-index, inverse-probability-of-censoring-weighted five-year area under the curve (AUC) and Brier score, calibration slope, and calibration-in-the-large. Internal validation used 500 bootstrap replicates; temporal validation trained the model in diagnoses through 2014 and evaluated it in diagnoses from 2015 onward. The model was benchmarked against PREDICT v2.2 surgery-only predictions in eligible nonmetastatic temporal-validation patients."),
        ("Results: ", "Among 4,231 patients, 492 deaths occurred during follow-up and 293 occurred within five years. Median follow-up was 94.1 months. After optimism correction, the five-year overall-survival model had a C-index of 0.749, AUC of 0.775, Brier score of 0.061, and calibration slope of 0.970. In temporal validation (n=1,654), the C-index was 0.762 and AUC was 0.779. In the PREDICT comparison (n=1,371; 77 five-year deaths), GBCS and PREDICT had similar AUCs (0.752 vs 0.758; difference -0.006, 95% bootstrap CI -0.035 to 0.027), whereas the GBCS Brier score was lower (0.0527 vs 0.0752). Mean five-year mortality was 6.46% with GBCS, 17.16% with PREDICT surgery only, and 6.05% observed."),
        ("Conclusions: ", "A six-predictor GBCS model provided useful discrimination and close five-year calibration within the cohort system. The PREDICT comparison supports similar ranking performance but cannot establish calibration superiority because treatment inputs were unavailable. Independent external validation, recalibration where needed, and prospective impact evaluation are required before clinical use."),
    ]
    for label, text in abstract_parts:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(label)
        font_run(r, bold=True)
        r = p.add_run(text)
        font_run(r)

    add_text(doc, "Keywords: breast cancer; prognosis; prediction model; overall survival; calibration; China; PREDICT", bold_prefix="Keywords: ", align=WD_ALIGN_PARAGRAPH.LEFT, after=12)

    h(doc, "Introduction", 1)
    add_text(doc, "Breast cancer is a major and growing health burden in China, with substantial variation in stage at diagnosis, tumor biology, treatment access, and outcomes across settings.1,2 Individualized estimates of absolute mortality risk can support prognosis discussions, risk-stratified follow-up, and the design of clinical studies, provided that the model is well calibrated in the population where it will be applied.")
    add_text(doc, "Several prognostic tools are available. PREDICT was developed using United Kingdom registry data to estimate survival after surgery and the absolute benefits of adjuvant treatment, and subsequent versions incorporated HER2 and contemporary treatment effects.5-7 Chinese cohorts have also produced conventional nomograms and machine-learning models, but many tools rely on treatment variables measured after diagnosis, restrict the target population, or have limited validation across time and clinical settings.3,4")
    add_text(doc, "The Guangzhou Breast Cancer Study (GBCS) is a prospective patient cohort established in 2008 across three hospitals in South China, with systematic collection of clinicopathological factors and longitudinal outcomes.1 We aimed to develop a parsimonious diagnosis-time model for five-year all-cause mortality using routinely available clinical predictors, evaluate optimism-corrected and temporal performance, and benchmark ranking and calibration against PREDICT v2.2. Five-year progression-free survival (PFS) was evaluated as a secondary outcome.")

    h(doc, "Methods", 1)
    h(doc, "Study design and participants", 2)
    add_text(doc, "This prognostic model development and internal validation study used the 2023 linked follow-up database of the GBCS. The parent cohort has been described previously.1 Eligible participants had invasive breast cancer, a diagnosis date from 1 October 2008 through 31 January 2018, nonnegative recorded follow-up times, and valid event indicators for overall survival (OS) and PFS. The prediction time origin was the date of diagnosis. The intended initial use is research and prognosis stratification rather than direct treatment selection.")

    h(doc, "Outcomes", 2)
    add_text(doc, "The primary outcome was all-cause mortality within five years after diagnosis. The secondary outcome was progression or death within five years, using the PFS definition in the source follow-up database. Participants without an event were censored at their last recorded follow-up. Follow-up duration was summarized using the reverse Kaplan-Meier method.")

    h(doc, "Candidate predictors", 2)
    add_text(doc, "The prespecified clinical model included age at diagnosis, AJCC stage (I-IV), estrogen receptor (ER), progesterone receptor (PR), HER2, and Ki-67. These variables were chosen because they are routinely available at diagnosis and have established prognostic relevance. Age was modeled continuously using a natural cubic spline with internal knots at 43 and 52 years and boundary knots at 19 and 97 years. No univariable screening was used, and all clinical predictors were retained. Treatment variables were excluded because they occurred after the prediction time and could introduce information leakage and treatment-policy dependence.")

    h(doc, "Missing data", 2)
    add_text(doc, "Predictor missingness ranged from 0.2% for age to 15.3% for Ki-67. Ten multiply imputed datasets were generated by chained equations with 10 iterations. Event indicators and Nelson-Aalen cumulative-hazard estimates were included in the imputation process. Model estimation and prediction were repeated within each imputed dataset, with coefficient uncertainty pooled using Rubin rules where applicable.")

    h(doc, "Model development and selection", 2)
    add_text(doc, "Separate Cox proportional-hazards models were fitted for OS and PFS. The primary clinical model was compared with an extended Cox model adding body mass index, menopausal status, education, parity, breastfeeding, and family history, and with an elastic-net Cox model using the extended predictor set. The final model was selected using validated discrimination, prediction error, calibration, parsimony, and feasibility of implementation. Individual five-year risks from the final model were averaged across the 10 imputation-specific clinical Cox fits.")

    h(doc, "Model performance and validation", 2)
    add_text(doc, "Discrimination was quantified using Harrell's C-index and a censoring-adjusted five-year AUC. Overall prediction error was measured using the inverse-probability-of-censoring-weighted Brier score. Calibration was assessed with calibration-in-the-large, the calibration slope, and observed versus mean predicted risk across tenths of predicted risk. Internal validation used 50 bootstrap samples within each of 10 imputations (500 replicates per model and endpoint) to estimate optimism. Temporal validation fitted models using patients diagnosed through 2014 and evaluated them among patients diagnosed from 2015 onward. Leave-one-hospital-out analyses and decision-curve analyses were prespecified as supporting analyses.")

    h(doc, "Comparison with PREDICT", 2)
    add_text(doc, "PREDICT v2.2 was used as a benchmark in temporal-validation patients with nonmetastatic disease, age 25-85 years, known ER status, and recorded T and N categories. Exact tumor size and positive-node count were used when available; otherwise, prespecified representative values were assigned for T1-T4 and N0-N3 categories. Detection method was set to PREDICT's unknown category. The GBCS Ki-67 threshold of at least 14% was mapped to the binary PREDICT input, which uses a threshold above 10%. Because chemotherapy generation, endocrine-therapy duration, trastuzumab exposure, and bisphosphonate treatment were not sufficiently complete, the comparison used PREDICT surgery-only survival. Model differences in AUC and Brier score were estimated using 500 paired bootstrap samples.")

    h(doc, "Statistical analysis and reporting", 2)
    add_text(doc, "Analyses were performed in R using survival, mice, and glmnet. Two-sided P values were descriptive because the study objective was prediction rather than causal inference. The manuscript was structured according to TRIPOD+AI, and model limitations were considered using PROBAST+AI principles.8,9 Modern sample-size guidance emphasizes participants, outcome events, and predictor parameters rather than a fixed events-per-variable rule.10,11")

    h(doc, "Ethics", 2)
    add_text(doc, "The Guangzhou Medical Ethics Committee of the Chinese Medical Association approved the parent GBCS (ID 2012-8), and all participants provided written informed consent.1 Any additional approval or waiver specific to this secondary prognostic-model analysis should be confirmed before submission.")

    h(doc, "Results", 1)
    h(doc, "Study population", 2)
    add_text(doc, "Of 5,412 records in the linked GBCS database, 355 were excluded because the classification did not indicate invasive disease, 701 because the diagnosis date was missing or outside the prespecified recruitment window, and 125 because survival follow-up was missing or invalid. The final analysis included 4,231 women with invasive breast cancer (Figure 1). Mean age at diagnosis was 47.9 years (SD 10.6), and median age was 47 years (IQR 41-55). Among patients with recorded stage, stage II was most frequent. During follow-up, 492 deaths and 810 PFS events occurred; 293 deaths and 570 PFS events occurred within five years. Median follow-up for OS was 94.1 months (7.84 years; IQR 73.2-129.4). Detailed cohort characteristics and missingness are provided in Supplementary Tables S1 and S7.")

    doc.add_page_break()
    add_figure(
        doc, FLOW_FIGURE, 6.2,
        "Figure 1", "Flow of participant selection, analysis sets, model development, and validation. Bootstrap, temporal, and leave-one-hospital-out procedures are internal validation within the GBCS system. The PREDICT analysis is a benchmark comparison, not external validation. ER, estrogen receptor; GBCS, Guangzhou Breast Cancer Study; HER2, human epidermal growth factor receptor 2; OS, overall survival; PFS, progression-free survival; PR, progesterone receptor."
    )

    h(doc, "Model selection and prognostic structure", 2)
    add_text(doc, "The final model retained age, stage, ER, PR, HER2, and Ki-67. The extended clinical model produced negligible optimism-corrected gains over the six-predictor model (five-year OS AUC 0.778 vs 0.775; Brier score 0.0611 vs 0.0610), whereas the one-standard-error elastic-net model underfit (AUC 0.719; calibration slope 2.64). The parsimonious clinical Cox model was therefore selected. Stage was the dominant prognostic factor: compared with stage I, the pooled hazard ratios for mortality were 1.57 (95% CI 1.17-2.11), 5.09 (3.72-6.97), and 15.83 (10.74-23.33) for stages II, III, and IV, respectively. PR positivity was associated with lower mortality, and Ki-67 at least 14% with higher mortality (Supplementary Table S2). These coefficients describe prognosis and should not be interpreted causally.")

    h(doc, "Internal and temporal validation", 2)
    add_text(doc, "After bootstrap optimism correction, the five-year OS model had a C-index of 0.749, AUC of 0.775, Brier score of 0.061, and calibration slope of 0.970. In temporal validation among 1,654 patients diagnosed from 2015 onward, the C-index was 0.762, AUC was 0.779, Brier score was 0.061, and calibration slope was 0.964 (Table 1). Calibration-in-the-large was close to zero in both analyses. The secondary PFS model showed lower discrimination but similar calibration slopes (Supplementary Tables S3 and S4; Supplementary Figures S1 and S2).")

    perf = pd.read_csv(PAPER / "table3_five_year_performance.csv")
    os_perf = perf[perf.endpoint == "OS"].copy()
    table1_rows = []
    for _, row in os_perf.iterrows():
        table1_rows.append([
            row["validation"], fmt(row["c_index"]), fmt(row["auc_5year"]),
            fmt(row["brier_5year"]), fmt(row["calibration_slope"]),
            f"{row['calibration_in_large'] * 100:.2f}%",
        ])
    add_caption(doc, "Table 1", "Performance of the final GBCS model for five-year all-cause mortality.")
    add_table(doc, ["Validation", "C-index", "5-y AUC", "5-y Brier", "Calibration slope", "Calibration-in-the-large"], table1_rows,
              [3000, 1100, 1100, 1100, 1450, 1610], numeric_cols={1, 2, 3, 4, 5}, font_size=8.5)
    add_text(doc, "AUC, area under the time-dependent receiver operating characteristic curve; GBCS, Guangzhou Breast Cancer Study. Calibration-in-the-large is mean predicted risk minus Kaplan-Meier observed risk.", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=10)

    h(doc, "Benchmark against PREDICT", 2)
    add_text(doc, "The temporal PREDICT comparison included 1,371 nonmetastatic patients and 77 deaths within five years. Discrimination was similar for GBCS and PREDICT v2.2 surgery-only predictions: AUCs were 0.752 and 0.758, respectively, with a paired difference of -0.006 (95% bootstrap CI -0.035 to 0.027). The Brier score was lower for GBCS (0.0527 vs 0.0752; paired difference -0.0225, 95% CI -0.0293 to -0.0151). Mean predicted mortality was 6.46% for GBCS and 17.16% for PREDICT surgery only, compared with Kaplan-Meier observed mortality of 6.05% (Table 2 and Figure 2). Across common GBCS-risk deciles, GBCS predictions followed the observed mortality gradient more closely, while PREDICT surgery-only risks were progressively higher in the upper deciles.")

    comp = pd.read_csv(PAPER / "table4_gbcs_predict_comparison.csv")
    table2_rows = []
    labels = {"GBCS_clinical_Cox": "GBCS clinical Cox", "PREDICT_v2.2_surgery_only": "PREDICT v2.2 surgery only"}
    for _, row in comp.iterrows():
        table2_rows.append([
            labels[row.model], fmt(row.auc_5year), f"{row.brier_5year:.4f}",
            fmt(row.calibration_slope), f"{100*row.mean_predicted_risk:.2f}%", f"{100*row.observed_risk:.2f}%",
        ])
    add_caption(doc, "Table 2", "Five-year mortality prediction in the temporal PREDICT-comparison cohort (n=1,371).")
    add_table(doc, ["Model", "5-y AUC", "5-y Brier", "Calibration slope", "Mean predicted risk", "Observed risk"], table2_rows,
              [2800, 1050, 1100, 1350, 1600, 1460], numeric_cols={1, 2, 3, 4, 5}, font_size=8.5)
    add_text(doc, "Paired AUC difference (GBCS minus PREDICT), -0.006 (95% bootstrap CI -0.035 to 0.027); paired Brier-score difference, -0.0225 (95% CI -0.0293 to -0.0151). PREDICT calibration used surgery-only survival because complete adjuvant-treatment inputs were unavailable.", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=10)

    add_figure(
        doc, PAPER / "figure1_gbcs_predict_observed_by_decile.png", 6.2,
        "Figure 2", "Observed and predicted five-year mortality by decile of GBCS-predicted risk. Gray bars show Kaplan-Meier observed mortality; lines show mean GBCS and PREDICT v2.2 surgery-only predictions. Deciles contain the same patients for all three series."
    )

    h(doc, "Discussion", 1)
    add_text(doc, "In this multicenter prospective cohort from South China, a diagnosis-time Cox model using six routinely available predictors achieved useful discrimination and close five-year calibration within the GBCS system. Performance remained similar when the model was trained in earlier diagnoses and evaluated in patients diagnosed from 2015 onward. The extended model offered no meaningful improvement, supporting selection of the more parsimonious model. A secondary PFS model was less discriminating, and hospital hold-out analyses suggested heterogeneity in PFS transportability.")
    add_text(doc, "The GBCS model and PREDICT ranked five-year mortality risk similarly in the temporal comparison, but their absolute predictions differed. PREDICT was originally developed for early invasive breast cancer after surgery in the United Kingdom and was designed to estimate both prognosis and treatment benefit.5-7 Our benchmark necessarily used surgery-only survival because detailed adjuvant-treatment inputs were incomplete. The resulting overprediction in a treated cohort is therefore expected and should not be interpreted as evidence that a fully specified contemporary PREDICT model is poorly calibrated in South China.")
    add_text(doc, "The model's five-year AUC of 0.775 after optimism correction is comparable with performance reported for clinical prognostic tools in Chinese breast cancer cohorts, although direct comparisons are limited by differences in populations, outcomes, predictor availability, and validation design.3,4 The present model prioritizes prediction at diagnosis and excludes treatment variables, which makes the time origin explicit and avoids leakage but means that predictions reflect the treatment patterns embedded in the development cohort. Calibration is particularly important for this use because good discrimination alone does not ensure reliable absolute risk estimates.12")
    add_text(doc, "Strengths include the prospective cohort design, recruitment from three clinical centers, nearly eight years of median follow-up, prespecified routinely available predictors, multiple imputation, and validation addressing optimism and temporal transportability. We also report calibration, prediction error, and a paired PREDICT benchmark rather than relying only on discrimination.")
    add_text(doc, "Several limitations require emphasis. First, bootstrap, temporal, and hospital hold-out analyses remain internal validation; the model has not been evaluated in an independent healthcare system. Second, proportional-hazards diagnostics were significant for the global OS and PFS models, particularly for age and hormone-receptor variables. Five-year calibration was acceptable, but regression coefficients should not be interpreted as constant causal effects, and flexible time-varying-effect sensitivity analyses are needed before submission. Third, missing predictor values were imputed, and predictor and outcome definitions depend on the source database. Fourth, the model predicts all-cause mortality and does not separate breast cancer deaths from competing mortality. Fifth, the PREDICT input mapping approximated tumor size and node count for some patients, used a mismatched Ki-67 cutoff, and lacked treatment information. Finally, treatment patterns and diagnostic practices have evolved since cohort recruitment, so future application will require contemporary external validation and, where necessary, recalibration.")
    add_text(doc, "The current model should therefore be viewed as a research tool for risk stratification and as a candidate for independent validation. The next steps are to evaluate transportability in a non-GBCS Chinese cohort, examine flexible survival models with time-varying effects, quantify performance in clinically important subgroups, and assess whether use of the model improves decisions or patient outcomes.")

    h(doc, "Conclusions", 1)
    add_text(doc, "A parsimonious GBCS model based on age, stage, ER, PR, HER2, and Ki-67 predicted five-year all-cause mortality with useful discrimination and close calibration during internal and temporal validation. Its discrimination was similar to PREDICT v2.2 in an eligible temporal subset, but calibration comparisons were constrained by unavailable treatment inputs. Independent external validation and impact evaluation are required before clinical implementation.")

    h(doc, "Declarations", 1)
    h(doc, "Funding", 2)
    add_text(doc, "To be completed by the authors using the exact grant names and award numbers supporting the GBCS and this analysis.")
    h(doc, "Author contributions", 2)
    add_text(doc, "To be completed after the author list and contribution roles have been agreed. CRediT roles are recommended.")
    h(doc, "Competing interests", 2)
    add_text(doc, "To be confirmed by all authors before submission.")
    h(doc, "Data and code availability", 2)
    add_text(doc, "Individual-level GBCS data are governed by cohort approvals and are not publicly released. De-identified data access may be considered following a methodologically sound proposal and required approvals. Analysis code, the de-identified model bundle, and scoring code can be shared with the publication subject to cohort governance.")

    h(doc, "References", 1)
    references = [
        "Wang J, Li N, Xiao CK, et al. Cohort profile: Guangzhou breast cancer study (GBCS). Eur J Epidemiol. 2024;39(12):1401-1410. doi:10.1007/s10654-024-01180-y.",
        "Fan L, Strasser-Weippl K, Li JJ, et al. Breast cancer in China. Lancet Oncol. 2014;15(7):e279-e289. doi:10.1016/S1470-2045(13)70567-9.",
        "Wang X, Feng Z, Huang Y, et al. A nomogram to predict the overall survival of breast cancer patients and guide postoperative adjuvant chemotherapy in China. Cancer Manag Res. 2019;11:10029-10039. doi:10.2147/CMAR.S215000.",
        "Zhong X, Luo T, Deng L, et al. Multidimensional machine learning personalized prognostic model in an early invasive breast cancer population-based cohort in China: algorithm validation study. JMIR Med Inform. 2020;8(11):e19069. doi:10.2196/19069.",
        "Wishart GC, Azzato EM, Greenberg DC, et al. PREDICT: a new UK prognostic model that predicts survival following surgery for invasive breast cancer. Breast Cancer Res. 2010;12(1):R1. doi:10.1186/bcr2464.",
        "Wishart GC, Bajdik CD, Dicks E, et al. PREDICT Plus: development and validation of a prognostic model for early breast cancer that includes HER2. Br J Cancer. 2012;107(5):800-807. doi:10.1038/bjc.2012.338.",
        "Grootes I, Wishart GC, Pharoah PDP. An updated PREDICT breast cancer prognostic model including the benefits and harms of radiotherapy. NPJ Breast Cancer. 2024;10(1):6. doi:10.1038/s41523-024-00612-y.",
        "Collins GS, Moons KGM, Dhiman P, et al. TRIPOD+AI statement: updated guidance for reporting clinical prediction models that use regression or machine learning methods. BMJ. 2024;385:e078378. doi:10.1136/bmj-2023-078378.",
        "Moons KGM, Damen JAA, Kaul T, et al. PROBAST+AI: an updated quality, risk of bias, and applicability assessment tool for prediction models using regression or artificial intelligence methods. BMJ. 2025;388:e082505. doi:10.1136/bmj-2024-082505.",
        "Riley RD, Snell KI, Ensor J, et al. Minimum sample size for developing a multivariable prediction model: Part II - binary and time-to-event outcomes. Stat Med. 2019;38(7):1276-1296. doi:10.1002/sim.7992.",
        "Riley RD, Ensor J, Snell KIE, et al. Calculating the sample size required for developing a clinical prediction model. BMJ. 2020;368:m441. doi:10.1136/bmj.m441.",
        "Van Calster B, McLernon DJ, van Smeden M, Wynants L, Steyerberg EW. Calibration: the Achilles heel of predictive analytics. BMC Med. 2019;17(1):230. doi:10.1186/s12916-019-1466-7.",
    ]
    for i, ref in enumerate(references, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.3)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(f"{i}. {ref}")
        font_run(r, size=9.5)

    # Supplementary material in the same editable file for easy submission splitting.
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("SUPPLEMENTARY MATERIALS")
    font_run(r, size=18, bold=True, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("GBCS five-year breast cancer prognosis prediction model")
    font_run(r, size=12, italic=True, color=TEXT_GRAY)

    h(doc, "Supplementary methods", 1)
    h(doc, "Prediction equation and implementation", 2)
    add_text(doc, "For imputation-specific model j, the five-year event risk for patient i was calculated as p_ij = 1 - exp[-H0_j(60) x exp(LP_ij)], where H0_j(60) is the uncentered baseline cumulative hazard at 60 months and LP_ij is the model linear predictor. The deployed prediction is the arithmetic mean of p_ij across the 10 imputation-specific clinical Cox models. The saved R model bundle contains every coefficient vector and baseline-hazard function; the accompanying scoring script applies the same factor coding, spline knots, and averaging procedure.")
    h(doc, "Model comparison", 2)
    add_text(doc, "The extended Cox model added BMI, menopausal status, education, parity, breastfeeding, and family history. The elastic-net model used alpha=0.5 with 10-fold cross-validation and the one-standard-error penalty. The clinical model was selected because the extended model's validated performance gain was negligible and the elastic-net model showed substantial underfitting.")

    # Supplementary Table S1
    t1 = pd.read_csv(PAPER / "table1_cohort_characteristics.csv")
    s1_rows = [[r.characteristic, "" if pd.isna(r.level) else r.level, r.value, str(int(r.missing))] for _, r in t1.iterrows()]
    add_caption(doc, "Supplementary Table S1", "Characteristics of the 4,231-patient analysis cohort.")
    add_table(doc, ["Characteristic", "Level", "Value", "Missing, n"], s1_rows,
              [3100, 2400, 2400, 1460], numeric_cols={2, 3}, font_size=8.5)
    add_text(doc, "Percentages use the full analysis cohort as denominator; category percentages therefore do not sum to 100% when data are missing.", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=10)

    # Supplementary Table S2
    t2 = pd.read_csv(PAPER / "table2_final_model_coefficients.csv")
    s2_rows = [[r.endpoint, r.term_label, f"{r.coefficient:.3f}", f"{r.standard_error:.3f}", r.hr_95ci, r.p_formatted] for _, r in t2.iterrows()]
    add_caption(doc, "Supplementary Table S2", "Pooled coefficients for the final clinical Cox models.")
    add_table(doc, ["Endpoint", "Predictor", "Beta", "SE", "Hazard ratio (95% CI)", "P value"], s2_rows,
              [850, 3650, 900, 900, 2000, 1060], numeric_cols={0, 2, 3, 4, 5}, font_size=7.8)
    add_text(doc, "Age was represented by a three-degree-of-freedom natural spline; individual spline coefficients should not be interpreted in isolation. Reference groups were stage I, ER negative, PR negative, HER2 negative, and Ki-67 <14%.", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=10)

    # Supplementary Table S3
    s3_rows = []
    for _, r in perf.iterrows():
        s3_rows.append([r.endpoint, r.validation, fmt(r.c_index), fmt(r.auc_5year), fmt(r.brier_5year), fmt(r.calibration_slope), f"{100*r.calibration_in_large:.2f}%"])
    add_caption(doc, "Supplementary Table S3", "Five-year performance of the final OS and PFS models.")
    add_table(doc, ["Endpoint", "Validation", "C-index", "AUC", "Brier", "Calibration slope", "Calibration-in-the-large"], s3_rows,
              [850, 2650, 900, 900, 900, 1400, 1760], numeric_cols={0, 2, 3, 4, 5, 6}, font_size=7.8)

    # Supplementary Table S4
    validated = pd.read_csv(RESULTS / "validated_model_performance.csv")
    validated = validated[validated.horizon_months == 60].copy()
    model_labels = {"clinical_cox": "Clinical Cox", "extended_cox": "Extended Cox", "elastic_net": "Elastic net"}
    s4_rows = []
    for _, r in validated.sort_values(["endpoint", "model"]).iterrows():
        s4_rows.append([r.endpoint, model_labels[r.model], fmt(r.corrected_c_index), fmt(r.corrected_auc), fmt(r.corrected_brier), fmt(r.validation_calibration_slope)])
    add_caption(doc, "Supplementary Table S4", "Bootstrap optimism-corrected comparison of candidate five-year models.")
    add_table(doc, ["Endpoint", "Model", "C-index", "AUC", "Brier", "Calibration slope"], s4_rows,
              [950, 2300, 1350, 1350, 1350, 2060], numeric_cols={0, 2, 3, 4, 5}, font_size=8.2)

    # Supplementary Table S5
    hospital = pd.read_csv(RESULTS / "hospital_validation_summary.csv")
    s5_rows = []
    for _, r in hospital.sort_values(["endpoint", "held_out_hospital"]).iterrows():
        s5_rows.append([r.endpoint, r.held_out_hospital, fmt(r.c_index), fmt(r.auc), fmt(r.brier), fmt(r.calibration_slope), f"{100*r.calibration_in_large:.2f}%"])
    add_caption(doc, "Supplementary Table S5", "Leave-one-hospital-out performance of the extended Cox model at five years.")
    add_table(doc, ["Endpoint", "Held-out hospital", "C-index", "AUC", "Brier", "Calibration slope", "Calibration-in-the-large"], s5_rows,
              [850, 2150, 950, 950, 950, 1450, 2060], numeric_cols={0, 2, 3, 4, 5, 6}, font_size=7.8)
    add_text(doc, "These hospital hold-outs are internal-external validation within the GBCS system and are not independent external validation.", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=10)

    # Supplementary Table S6
    s6_rows = [
        ["Population", "Diagnosis from 2015; M0; age 25-85 years; known ER; recorded T and N categories"],
        ["Tumor size", "Exact millimeters when available; otherwise T1=15, T2=35, T3=60, T4=50 mm"],
        ["Positive nodes", "Exact count when available; otherwise N0=0, N1=1, N2=4, N3=10"],
        ["Grade", "Recorded grade 1-3; PREDICT unknown/default handling when unavailable"],
        ["HER2", "Recorded negative, positive, or unknown"],
        ["Ki-67", "GBCS <14% vs >=14% mapped to PREDICT binary input (>10% threshold)"],
        ["Detection", "Unavailable; set to PREDICT unknown"],
        ["Treatment", "Surgery-only prediction because detailed systemic and radiation treatment inputs were incomplete"],
    ]
    add_caption(doc, "Supplementary Table S6", "Mapping of GBCS variables to PREDICT v2.2 inputs.")
    add_table(doc, ["Input", "Operational definition"], s6_rows, [2200, 7160], font_size=8.5)

    # Supplementary Table S7
    missing = pd.read_csv(RESULTS / "predictor_missingness.csv")
    keep = ["age", "stage", "er", "pr", "her2", "ki67", "bmi", "menopause", "education", "parity", "breastfeeding", "family_history"]
    missing = missing[missing.variable.isin(keep)].copy()
    variable_labels = {
        "age": "Age", "stage": "Stage", "er": "ER", "pr": "PR", "her2": "HER2", "ki67": "Ki-67",
        "bmi": "BMI", "menopause": "Menopausal status", "education": "Education", "parity": "Parity",
        "breastfeeding": "Breastfeeding", "family_history": "Family history",
    }
    s7_rows = [[variable_labels[r.variable], str(int(r.non_missing)), str(int(r.missing)), f"{r.missing_percent:.1f}%"] for _, r in missing.iterrows()]
    add_caption(doc, "Supplementary Table S7", "Predictor missingness before multiple imputation.")
    add_table(doc, ["Predictor", "Observed, n", "Missing, n", "Missing, %"], s7_rows, [3900, 1820, 1820, 1820], numeric_cols={1, 2, 3}, font_size=8.5)

    # Supplementary Table S8: PH diagnostic consistency.
    ph = pd.read_csv(RESULTS / "proportional_hazards_diagnostics.csv")
    ph["term_short"] = ph.term.str.replace(r"ns\(age.*", "Age spline", regex=True)
    ph["term_short"] = ph.term_short.replace({"stage": "Stage", "er": "ER", "pr": "PR", "her2": "HER2", "ki67": "Ki-67", "GLOBAL": "Global test"})
    ph_sum = ph.groupby(["endpoint", "term_short"], as_index=False).agg(
        imputations_p_lt_005=("p_value", lambda x: int((x < 0.05).sum())),
        median_p=("p_value", "median"),
    )
    term_order = {"Global test": 0, "Age spline": 1, "Stage": 2, "ER": 3, "PR": 4, "HER2": 5, "Ki-67": 6}
    ph_sum["ord"] = ph_sum.term_short.map(term_order)
    ph_sum = ph_sum.sort_values(["endpoint", "ord"])
    s8_rows = [[r.endpoint, r.term_short, f"{int(r.imputations_p_lt_005)}/10", "<0.001" if r.median_p < 0.001 else f"{r.median_p:.3f}"] for _, r in ph_sum.iterrows()]
    add_caption(doc, "Supplementary Table S8", "Proportional-hazards diagnostics across imputed datasets.")
    add_table(doc, ["Endpoint", "Term", "Imputations with P<0.05", "Median P value"], s8_rows, [1200, 3150, 2650, 2360], numeric_cols={0, 2, 3}, font_size=8.5)
    add_text(doc, "Schoenfeld-residual tests used the Kaplan-Meier transformation. These results motivate a flexible time-varying-effect sensitivity analysis before journal submission.", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, before=4, after=10)

    h(doc, "Supplementary figures", 1)
    add_figure(doc, PAPER / "figure2_five_year_calibration.png", 6.2, "Supplementary Figure S1", "Five-year calibration of the final GBCS OS and PFS models. Points represent tenths of predicted risk; the diagonal indicates perfect calibration.")
    add_figure(doc, PAPER / "figure3_validation_discrimination.png", 6.2, "Supplementary Figure S2", "Discrimination after bootstrap optimism correction and in temporal validation. AUC, area under the time-dependent receiver operating characteristic curve; OS, overall survival; PFS, progression-free survival.")
    add_figure(doc, RESULTS / "decision_curve_5year.png", 6.2, "Supplementary Figure S3", "Apparent five-year decision curves for candidate models. These analyses are exploratory and were not independently validated.")

    h(doc, "Supplementary files for reproducibility", 1)
    add_text(doc, "The project includes the analysis script, PREDICT comparison script, model-preparation script, de-identified model bundle, and scoring function. The full imputation-specific coefficient and baseline-hazard objects are stored in the model bundle rather than reproduced as a very large printed table.")

    # Core properties and save.
    doc.core_properties.title = "GBCS five-year breast cancer prognosis model manuscript"
    doc.core_properties.subject = "Development and internal validation of a five-year prognosis model"
    doc.core_properties.author = "GBCS manuscript team"
    doc.core_properties.keywords = "breast cancer; prediction model; prognosis; GBCS; PREDICT"
    doc.core_properties.comments = "Draft for author review; author list and declarations require confirmation."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_document()
    print(output)
