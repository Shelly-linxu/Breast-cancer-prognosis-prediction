"""Add the verified static GBCS calculator to the current manuscript.

The source manuscript is preserved. The output receives focused additions to
Methods, Discussion, Data and code availability, and Supplementary materials.
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


REPOSITORY = Path(__file__).resolve().parents[2]
WORKSPACE = REPOSITORY.parent
SOURCE = WORKSPACE / "GBCS_five_year_prognosis_model_manuscript v5 significance and references.docx"
OUTPUT = WORKSPACE / "GBCS_five_year_prognosis_model_manuscript v6 with calculator.docx"
SCREENSHOT = REPOSITORY / "calculator" / "calculator-preview.png"
PUBLIC_URL = "https://shelly-linxu.github.io/Breast-cancer-prognosis-prediction/calculator/"
REPOSITORY_URL = "https://github.com/Shelly-linxu/Breast-cancer-prognosis-prediction"


def find_paragraph(document, prefix):
    for paragraph in document.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def insert_before(anchor, text=""):
    paragraph = anchor.insert_paragraph_before(text)
    paragraph.style = anchor.style
    return paragraph


def append_after(anchor, paragraph):
    anchor._p.addnext(paragraph._p)


def new_paragraph_like(document, anchor, text=""):
    paragraph = document.add_paragraph(text)
    paragraph.style = anchor.style
    return paragraph


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not SCREENSHOT.exists():
        raise FileNotFoundError(SCREENSHOT)

    document = Document(SOURCE)

    statistical_heading = find_paragraph(document, "Statistical analysis and reporting")
    web_heading = insert_before(statistical_heading, "Web calculator")
    web_heading.runs[0].italic = True
    insert_before(
        statistical_heading,
        "We implemented the locked five-year OS model as an English-language static web calculator "
        "(GBCS 5-year OS Calculator, Research Preview v1.0). The tool requires age at diagnosis, "
        "AJCC overall stage at initial diagnosis, ER, PR, HER2, and Ki-67 and returns five-year "
        "all-cause mortality, five-year overall survival, and the range across the 10 imputed model "
        "fits. Browser calculations reproduce the imputation-specific Cox equations and baseline "
        "hazards before averaging predicted risks. All processing occurs locally in the browser; "
        "no patient information is transmitted, retained, or placed in the URL. Browser output was "
        "checked against the original R scorer at the age boundaries, spline knots, and all categorical "
        "levels. The calculator is labelled for research use only because independent external validation "
        "and clinical-impact evaluation remain necessary.",
    )

    discussion = find_paragraph(document, "Six variables routinely available")
    discussion.add_run(
        " To make the model transparent and immediately testable, we translated the locked equations "
        "into a browser-based calculator that can support research and structured prognosis discussions; "
        "its practical value will depend on independent external validation, prospective assessment of "
        "clinical impact, and recalibration as treatment patterns and baseline survival change."
    )

    availability = find_paragraph(document, "Individual-level GBCS data are governed")
    availability.text = (
        "Individual-level GBCS data are governed by cohort approvals and are not publicly released. "
        "De-identified data access may be considered following a methodologically sound proposal and "
        "required approvals. Analysis code, browser-calculator source, model-parameter export code, "
        f"and scoring tests are available at {REPOSITORY_URL}. The GBCS 5-year OS Calculator, Research "
        f"Preview v1.0, is available at {PUBLIC_URL}. The calculator contains model parameters only and "
        "does not contain participant-level data."
    )

    model_comparison_heading = find_paragraph(document, "Model comparison")
    implementation_heading = insert_before(model_comparison_heading, "Static web-calculator implementation")
    implementation_heading.runs[0].bold = True
    insert_before(
        model_comparison_heading,
        "The calculator applies the same age spline, factor coding, 60-month baseline cumulative hazards, "
        "and arithmetic averaging across the 10 imputation-specific OS models as the R scoring function. "
        "Inputs are restricted to the development ranges and categories. A prediction is not produced if "
        "age lies outside 19-97 years or if any required predictor is unavailable. The displayed model-fit "
        "range is the minimum to maximum prediction across the 10 imputed fits; it is not a 95% confidence "
        "interval and does not represent all uncertainty in an individual prediction."
    )
    guide_heading = insert_before(model_comparison_heading, "Calculator user guide")
    guide_heading.runs[0].bold = True
    insert_before(
        model_comparison_heading,
        "Users enter values recorded at initial diagnosis before treatment: age, AJCC overall stage I-IV, "
        "ER, PR, HER2 (negative, equivocal, or positive), and Ki-67 (<14% or at least 14%). Selecting "
        "Calculate prognosis displays the estimated probability of all-cause mortality within five years, "
        "the complementary five-year overall-survival probability, and the model-fit range. Results should "
        "not be described as low, intermediate, or high risk because clinical thresholds were not validated. "
        "The estimate reflects outcomes and treatment patterns observed in GBCS and must not be used alone "
        "to select treatment, alter follow-up, or replace clinical judgement."
    )

    last_caption = find_paragraph(document, "Supplementary Figure S3")

    # The retained source keeps the S2 caption on one over-wide line in the
    # headless renderer. A deliberate line break prevents edge clipping.
    s2_caption = find_paragraph(document, "Supplementary Figure S2")
    s2_caption.text = (
        "Supplementary Figure S2.\n"
        "Discrimination after bootstrap optimism correction and in temporal validation."
    )
    s2_caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in s2_caption.runs:
        run.bold = True

    figure_paragraph = new_paragraph_like(document, last_caption)
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_paragraph.paragraph_format.page_break_before = True
    figure_paragraph.add_run().add_picture(str(SCREENSHOT), width=Inches(6.35))
    append_after(last_caption, figure_paragraph)

    caption = new_paragraph_like(
        document,
        last_caption,
        "Supplementary Figure S4. GBCS 5-year OS Calculator research-preview interface. "
        "The example uses a synthetic profile and contains no participant-level information.",
    )
    append_after(figure_paragraph, caption)

    note = new_paragraph_like(
        document,
        last_caption,
        "Note: The calculator reports continuous five-year all-cause mortality and overall-survival "
        "estimates without clinical risk categories. External validation and clinical-impact evaluation "
        "are pending.",
    )
    append_after(caption, note)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
